"""Verträge der gemeinsamen Reporting-Pipeline für Schritt 10."""

import copy
from io import BytesIO
from numbers import Real
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook

import framework_mvp.application.modellausgabe_service as ausgabe_modul
import framework_mvp.reporting.xlsx_renderer as xlsx_modul
from framework_mvp.application.modellausgabe_service import ModellausgabeService
from framework_mvp.application.modellvalidierung_service import ModellvalidierungService
from framework_mvp.infrastructure.exceptions import Importintegritaetsfehler
from framework_mvp.reporting.asset_resolver import ReportAssetFehler, resolve_report_assets
from framework_mvp.reporting.html_renderer import render_report_html, template_verzeichnis
from framework_mvp.reporting.pdf_renderer import render_report_pdf
from framework_mvp.reporting.report_data import (
    ERWARTETE_BESTANDTEIL_IDS,
    ReportDataFehler,
    build_report_data,
)
from framework_mvp.reporting.xlsx_renderer import SHEET_NAMES, render_report_xlsx
from framework_mvp.workspace import WorkspaceKonfiguration


def _information(referenz: str, wert: object, artefakt: str = "E*") -> dict[str, object]:
    return {
        "informations_id": str(uuid4()),
        "strukturreferenz": referenz,
        "wert": wert,
        "herkunftsartefakt": artefakt,
        "herkunftsartefakt_id": str(uuid4()),
        "herkunftsartefakt_sha256": "a" * 64,
        "uebernahmeart": "metadatenzusammenfassung",
    }


def _k_stern(*, neue_felder: bool = True) -> dict[str, object]:
    projekt_id, analyse_id = uuid4(), uuid4()
    bestandteile = []
    for bestandteil_id in ERWARTETE_BESTANDTEIL_IDS:
        informationen: list[dict[str, object]] = []
        menschliche_eintraege: list[dict[str, object]] = []
        if bestandteil_id == "problemstellung":
            informationen.append(
                _information("untersuchungsauftrag.problemstellung", "Materialfluss prüfen", "U")
            )
        elif bestandteil_id == "ausgaben" and neue_felder:
            informationen.extend(
                [
                    _information(
                        "kpi_ergebnisse[0]",
                        {
                            "kpi_id": "mittlere_durchfuehrungszeit",
                            "bezeichnung": "Mittlere Durchführungszeit",
                            "status": "berechnet",
                            "ergebnis": 1850.1799999999998,
                            "einheit": "s",
                            "bezugsmenge": "Produktionsaufträge n",
                            "formel": "Σ Durchführungszeit_i / n",
                            "definitionsversion": 2,
                            "behandlungsart": "automatisch_berechnen",
                        },
                        "A_G",
                    ),
                    _information(
                        "kpi_ergebnisse[1]",
                        {
                            "kpi_id": "servicegrad",
                            "bezeichnung": "Servicegrad",
                            "status": "fuer_spaetere_manuelle_berechnung_vorgesehen",
                            "ergebnis": None,
                            "einheit": "%",
                            "bezugsmenge": "Kundenauftragspositionen",
                            "formel": "befriedigte Positionen / Positionen · 100",
                            "definitionsversion": 1,
                            "behandlungsart": "spaeter_manuell_berechnen",
                        },
                        "A_G",
                    ),
                    _information(
                        "conformance_checking",
                        {
                            "durchgefuehrt": True,
                            "ergebnis": {
                                "fitness": 0.9975345167,
                                "produzierte_tokens": 20,
                                "konsumierte_tokens": 19,
                                "fehlende_tokens": 1,
                                "verbleibende_tokens": 2,
                                "ausgewertete_faelle": 3,
                                "konforme_faelle": 2,
                                "abweichende_faelle": 1,
                            },
                        },
                        "A_G",
                    ),
                    _information(
                        "strukturierte_ergebnisse.performance_und_engpassanalyse",
                        {
                            "dt_db_ergebnis": {
                                "dt_statistik": {
                                    "anzahl": 3,
                                    "verspaetet": 1,
                                    "planmaessig": 1,
                                    "vorzeitig": 1,
                                    "mittelwert_sekunden": 20.0,
                                    "median_sekunden": 0.0,
                                },
                                "db_statistik": {
                                    "anzahl": 2,
                                    "laenger_als_geplant": 1,
                                    "gleich_geplant": 0,
                                    "kuerzer_als_geplant": 1,
                                    "mittelwert_sekunden": 10.0,
                                    "median_sekunden": 10.0,
                                },
                            },
                            "busy_ratio_ergebnis": {
                                "ressourcenstatistiken": [
                                    {
                                        "ressource": "M1",
                                        "anzahl_gueltige_busy_ratios": 2,
                                        "mittelwert_busy_ratio": 0.8,
                                        "median_busy_ratio": 0.8,
                                    }
                                ],
                                "potenzieller_engpass": "M1",
                            },
                        },
                        "A_G",
                    ),
                ]
            )
        elif bestandteil_id == "aktivitaeten":
            informationen.append(_information("sichtbare_aktivitaeten", ["A", "B"], "P"))
        elif bestandteil_id == "warteschlangen" and neue_felder:
            informationen.append(
                _information(
                    "strukturierte_ergebnisse.warteschlangen_und_wartezeiten",
                    {
                        "status": "ableitbar",
                        "berechnungsregel": "Start(B) − Ende(A)",
                        "potenzielle_wartezeiten": [
                            {
                                "von_aktivitaet": "A",
                                "zu_aktivitaet": "B",
                                "statistik": {
                                    "anzahl": 2,
                                    "mittelwert_sekunden": 90.0,
                                    "median_sekunden": 90.0,
                                },
                            }
                        ],
                    },
                    "A_G",
                )
            )
        elif bestandteil_id == "ressourcen" and neue_felder:
            informationen.append(
                _information(
                    "strukturierte_ergebnisse.ressourcen",
                    {
                        "modus": "manuell",
                        "herkunft": "menschlich bestätigte Zuordnung in Schritt 7",
                        "quellspalte": "",
                        "zuordnungen": [{"aktivitaet": "A", "ressourcen": ["M1", "M2"]}],
                    },
                    "A_G",
                )
            )
        elif bestandteil_id == "vereinfachungen" and neue_felder:
            informationen.append(
                _information(
                    "strukturierte_ergebnisse.vereinfachungen.etl_abstraktionen",
                    [
                        {
                            "quellspalte": "Von",
                            "vergleichsart": "Beginnt mit",
                            "suchwert_muster": "HRL-04-",
                            "vorher_muster": "HRL-04-*",
                            "abstraktionswert": "HRL-04",
                            "zielspalte": "Von_aggregiert",
                            "betroffene_beobachtungen": 185,
                            "originalwerte_erhalten": True,
                        }
                    ],
                    "A_G",
                )
            )
        elif bestandteil_id == "datenauswahl" and neue_felder:
            informationen.append(
                _information(
                    "strukturierte_ergebnisse.zeitbezogene_datenauswahl",
                    {
                        "bestaetigte_datenbasis": ["Q", "R", "T", "E*"],
                        "ankunftsregel": (
                            "Erster gültiger kanonischer Ereigniszeitstempel je Fall."
                        ),
                        "system_zwischenankunftszeit": {
                            "anzahl_entitaeten": 3,
                            "status": "ableitbar",
                            "statistik": {
                                "anzahl": 2,
                                "mittelwert_sekunden": 120.0,
                                "median_sekunden": 120.0,
                            },
                        },
                        "zwischenankunftszeiten": [
                            {
                                "definition": {"bezeichnung": "Auftragseingang"},
                                "statistik": {
                                    "anzahl": 2,
                                    "mittelwert_sekunden": 120.0,
                                    "median_sekunden": 120.0,
                                },
                            }
                        ],
                        "bearbeitungszeiten": [
                            {
                                "aktivitaet": "A",
                                "statistik": {
                                    "anzahl": 2,
                                    "mittelwert_sekunden": 60.0,
                                    "median_sekunden": 60.0,
                                },
                            }
                        ],
                        "vereinfachte_zeitspannen_bestaetigt": True,
                        "vereinfachungsentscheidung": (
                            "Mangels separatem Endzeitpunkt als vereinfachte Zeitspanne übernehmen."
                        ),
                        "vereinfachte_zeitspannen": [
                            {
                                "von_aktivitaet": "B",
                                "zu_aktivitaet": "C",
                                "statistik": {
                                    "anzahl": 1,
                                    "mittelwert_sekunden": 180.0,
                                    "median_sekunden": 180.0,
                                },
                            }
                        ],
                        "potenzielle_wartezeiten": [
                            {
                                "von_aktivitaet": "A",
                                "zu_aktivitaet": "B",
                                "statistik": {
                                    "anzahl": 2,
                                    "mittelwert_sekunden": 90.0,
                                    "median_sekunden": 90.0,
                                },
                            }
                        ],
                    },
                    "A_G",
                )
            )
            informationen.append(
                _information(
                    "strukturierte_ergebnisse.datenaufbereitung",
                    {
                        "ausgang": "ursprüngliche Datenquelle D",
                        "fachliche_bedeutung": (
                            "Alle Transformationen bilden eine Aufbereitungskette; "
                            "fachlich gültig ist genau der aktuelle Zwischendatensatz T."
                        ),
                        "transformationshistorie": [
                            {
                                "reihenfolge": 1,
                                "betroffener_datensatz": "Zwischendatensatz T",
                                "transformationsart": "Werte regelbasiert abstrahieren",
                                "betroffene_spalten": ["Von", "Zu"],
                                "regel": {"vergleichsart": "Beginnt mit", "suchwert": "HRL-04"},
                                "ersatz_oder_abstraktionswert": "HRL",
                                "beschreibung": "Lagerplätze fachlich gruppieren",
                                "eingang": "ursprüngliche Datenquelle D",
                                "ergebnis": "aktiver Zwischendatensatz T",
                                "wirkung": "389 Werte regelbasiert abstrahiert",
                                "zeilen_vorher": 400,
                                "zeilen_nachher": 400,
                            }
                        ],
                        "aktiver_zwischendatensatz_t": {
                            "id": str(uuid4()),
                            "zeilenanzahl": 400,
                            "spaltenanzahl": 8,
                            "sha256": "e" * 64,
                        },
                    },
                    "A_G",
                )
            )
        elif bestandteil_id == "darstellung_der_vorgaenge_des_systems":
            informationen.append(
                _information(
                    "prozessmodell_referenz",
                    {
                        "prozessmodell_id": str(analyse_id),
                        "process_mining_analyse_id": str(analyse_id),
                        "notation": "petrinetz",
                        "relativer_pfad": "modell.pnml",
                    },
                    "P",
                )
            )
        bestandteile.append(
            {
                "bestandteil_id": bestandteil_id,
                "bezeichnung": bestandteil_id.replace("_", " ").title(),
                "validierungsstatus": "fachlich_validiert",
                "urspruenglicher_bestandteil": {
                    "status": "vollstaendig_zugeordnet",
                    "verwendete_quellen": ["E*"],
                    "informationen": informationen,
                },
                "menschliche_eintraege": menschliche_eintraege,
            }
        )
    return {
        "artefaktart": "fachlich_validiertes_modell_k_stern",
        "artefaktversion": 1,
        "projekt_id": str(projekt_id),
        "k_stern_id": str(uuid4()),
        "validierungslauf_id": str(uuid4()),
        "erstellt_am": "2026-08-12T10:00:00+00:00",
        "modellbestandteile": bestandteile,
        "gesamtvalidierung": {
            "status": "fachlich_validiert",
            "validierungsvermerk": "Geprüft",
            "menschlich_bestaetigt": True,
        },
        "behandlungen_offener_eintraege": [],
        "k_referenz": {},
        "o_referenz": {},
        "eingabefingerabdruck": "b" * 64,
        "entscheidungsfingerabdruck": "c" * 64,
        "gesamtpruefsumme": "d" * 64,
    }


def test_build_report_data_projiziert_neue_felder_ohne_k_stern_mutation() -> None:
    k_stern = _k_stern()
    vorher = copy.deepcopy(k_stern)

    report = build_report_data(k_stern)

    assert k_stern == vorher
    assert report["warteschlangen"]["wartestellenhinweise"][0]["anzahl"] == 2
    assert report["ressourcen"]["aktivitaet_ressourcen"] == [
        {"aktivitaet": "A", "ressourcen": ["M1", "M2"]}
    ]
    assert report["ressourcen"]["zuordnungsmodus"] == "manuell"
    assert report["ressourcen"]["zuordnungsherkunft"].endswith("Schritt 7")
    assert report["ausgaben_und_eingaben"]["kpi_ergebnisse"][0]["ergebnis_anzeige"] == "1850,18 s"
    assert report["ausgaben_und_eingaben"]["kpi_ergebnisse"][1]["ergebnis_anzeige"] == (
        "Für spätere manuelle Berechnung vorgesehen"
    )
    assert (
        report["ausgaben_und_eingaben"]["conformance_checking"]["ergebnis"]["fitness"]
        == 0.9975345167
    )
    assert (
        report["ausgaben_und_eingaben"]["conformance_checking"]["ergebnis"]["fitness_anzeige"]
        == "99,75 %"
    )
    assert (
        report["ausgaben_und_eingaben"]["performance_und_engpassanalyse"]["dt_db_ergebnis"][
            "dt_statistik"
        ]["anzahl"]
        == 3
    )
    assert (
        report["daten"]["zeitbezogene_datenauswahl"]["zwischenankunftszeiten"][0]["statistik"][
            "median_sekunden"
        ]
        == 120.0
    )
    assert (
        report["daten"]["zeitbezogene_datenauswahl"]["system_zwischenankunftszeit"]["statistik"][
            "anzahl"
        ]
        == 2
    )
    assert (
        report["daten"]["zeitbezogene_datenauswahl"]["vereinfachte_zeitspannen"][0]["statistik"][
            "median_sekunden"
        ]
        == 180.0
    )
    assert report["vereinfachungen"]["vereinfachte_zeitspannen"] == {
        "status": "Menschlich bestätigt",
        "entscheidung": ("Mangels separatem Endzeitpunkt als vereinfachte Zeitspanne übernehmen."),
        "betroffene_uebergaenge": 1,
        "fachliche_grenze": (
            "Gemeinsame Zeitspanne aus Bearbeitung, Transport, Warten und sonstigen "
            "Zwischenzeiten; keine zusätzliche Bearbeitungs- oder Wartezeit für "
            "denselben Abschnitt."
        ),
    }
    assert (
        report["daten"]["datenaufbereitung"]["transformationshistorie"][0]["ergebnis"]
        == "aktiver Zwischendatensatz T"
    )
    assert len(report["modellbestandteile"]) == 16


def test_reportgliederung_folgt_den_acht_fachlichen_abschnitten() -> None:
    html = render_report_html(build_report_data(_k_stern()))

    for nummer, titel in (
        ("1", "Problemstellung"),
        ("2", "Zielsetzung"),
        ("3", "Ausgaben und Eingaben"),
        ("4", "Modellumfang, Modellgrenzen und Detaillierungsgrad"),
        ("4.1", "Entitäten"),
        ("4.2", "Aktivitäten"),
        ("4.3", "Warteschlangen"),
        ("4.4", "Ressourcen"),
        ("5", "Annahmen und Vereinfachungen"),
        ("6", "Datenauswahl und Daten"),
        ("7", "Darstellung der Vorgänge des Systems"),
        ("8", "Technische Nachvollziehbarkeit"),
    ):
        assert f'<div class="component-number">\n        {nummer}\n' in html or (
            f'<div class="section-number">{nummer}</div>' in html
        )
        assert titel in html
    assert "1850.1799999999998" not in html
    assert "1850,18 s" in html
    assert "99,75 %" in html
    assert "0.9975345167" not in html
    assert "12.08.2026 10:00:00 +00:00" in html


@pytest.mark.parametrize(
    (
        "fallstudie",
        "systemtyp",
        "kpi_id",
        "kpi_name",
        "ergebnis",
        "einheit",
        "erwartete_anzeige",
    ),
    (
        (
            "intralogistiksystem",
            "intralogistik",
            "servicegrad",
            "Servicegrad",
            3.333,
            "%",
            "3,33 %",
        ),
        (
            "produktionssystem",
            "produktion",
            "einhaltung_lagerbandbreite",
            "Einhaltung Lagerbandbreite",
            90.0,
            "%",
            "90 %",
        ),
        (
            "synthetisches-produktionssystem",
            "produktion",
            "mittlere_durchfuehrungszeit",
            "Mittlere Durchführungszeit",
            1850.1799999999998,
            "s",
            "1850,18 s",
        ),
    ),
)
def test_reports_der_drei_fallstudien_bleiben_renderbar(
    tmp_path: Path,
    fallstudie: str,
    systemtyp: str,
    kpi_id: str,
    kpi_name: str,
    ergebnis: float,
    einheit: str,
    erwartete_anzeige: str,
) -> None:
    k_stern = _k_stern()
    bestandteile = cast(list[dict[str, Any]], k_stern["modellbestandteile"])
    umfang = next(wert for wert in bestandteile if wert["bestandteil_id"] == "modellumfang")
    umfang["urspruenglicher_bestandteil"]["informationen"].append(
        _information("systemprofil", {"systemtyp": systemtyp}, "S")
    )
    ausgaben = next(wert for wert in bestandteile if wert["bestandteil_id"] == "ausgaben")
    kpi = ausgaben["urspruenglicher_bestandteil"]["informationen"][0]["wert"]
    kpi.update(
        {
            "kpi_id": kpi_id,
            "bezeichnung": kpi_name,
            "ergebnis": ergebnis,
            "einheit": einheit,
        }
    )

    report = build_report_data(k_stern, projektbezeichnung=fallstudie)
    html = render_report_html(report)
    pdf = render_report_pdf(report, tmp_path / f"{fallstudie}.pdf")

    assert report["modellumfang"]["systemtyp"] == systemtyp
    assert kpi_name in html
    assert erwartete_anzeige in html
    assert str(ergebnis) not in html
    assert pdf.read_bytes().startswith(b"%PDF-")


def test_transformationshistorie_wird_einmal_fachlich_in_html_und_pdf_ausgegeben(
    tmp_path: Path,
) -> None:
    k_stern = _k_stern()
    report = build_report_data(k_stern)

    historie = report["daten"]["datenaufbereitung"]["transformationshistorie"]
    assert len(historie) == 1
    assert historie[0]["betroffene_spalten"] == ["Von", "Zu"]
    assert historie[0]["ersatz_oder_abstraktionswert"] == "HRL"
    html = render_report_html(report)
    assert html.count("Datenaufbereitung: D → aktiver Zwischendatensatz T") == 1
    assert "Regelbasierte ETL-Abstraktionen" not in html
    assert "Werte regelbasiert abstrahieren" in html
    assert "HRL-04" in html
    assert "Von, Zu" in html
    assert "389 Werte regelbasiert abstrahiert" in html

    ziel = render_report_pdf(report, tmp_path / "abstraktionen.pdf")
    assert ziel.read_bytes().startswith(b"%PDF-")


def test_report_nutzt_potenzielle_wartezeiten_aus_datenauswahl_ohne_warteschlange() -> None:
    k_stern = _k_stern()
    bestandteile = cast(list[dict[str, Any]], k_stern["modellbestandteile"])
    warteschlangen = next(
        wert for wert in bestandteile if wert["bestandteil_id"] == "warteschlangen"
    )
    warteschlangen["urspruenglicher_bestandteil"]["informationen"] = []

    report = build_report_data(k_stern)

    assert report["warteschlangen"]["wartestellenhinweise"] == [
        {
            "uebergang": {"von": "A", "zu": "B"},
            "anzahl": 2,
            "mittlere_wartezeit_sekunden": 90.0,
            "mediane_wartezeit_sekunden": 90.0,
        }
    ]


def test_xlsx_renderer_erzeugt_zehn_geordnete_lesbare_arbeitsblaetter() -> None:
    report = build_report_data(
        _k_stern(),
        projektbezeichnung="Fördertechnik Süd",
        softwareversion="0.1.0-test",
    )

    inhalt = render_report_xlsx(report)
    arbeitsmappe = load_workbook(BytesIO(inhalt), data_only=False)

    assert arbeitsmappe.sheetnames == list(SHEET_NAMES)
    assert all(not blatt.sheet_view.showGridLines for blatt in arbeitsmappe.worksheets)
    assert arbeitsmappe["Übersicht"]["B6"].value == "Fördertechnik Süd"
    assert any(
        zelle.value == "Geprüft"
        for zeile in arbeitsmappe["Validierung"].iter_rows()
        for zelle in zeile
    )
    assert not any(
        isinstance(zelle.value, str) and zelle.value.startswith("=")
        for blatt in arbeitsmappe.worksheets
        for zeile in blatt.iter_rows()
        for zelle in zeile
    )
    assert all(
        zelle.font.name == "Calibri"
        for blatt in arbeitsmappe.worksheets
        for zeile in blatt.iter_rows()
        for zelle in zeile
        if zelle.value is not None
    )
    datenblatt = arbeitsmappe["Daten & Datenauswahl"]
    kopfzeile = next(
        zeile
        for zeile in datenblatt.iter_rows()
        if zeile[0].value == "Kennwert" and zeile[1].value == "Bezug"
    )
    kopfwerte = [zelle.value for zelle in kopfzeile if zelle.value is not None]
    assert "Minimum" not in kopfwerte and "Maximum" not in kopfwerte
    statistikwerte = [
        zeile
        for zeile in datenblatt.iter_rows(
            min_row=cast(int, kopfzeile[0].row) + 1,
            values_only=False,
        )
        if zeile[0].value
        in {
            "System-IAT nach Gl. 3.16",
            "Zwischenankunftszeit",
            "Bearbeitungszeit",
            "Vereinfachte Start-zu-Start-Zeitspanne",
            "Wartezeit",
        }
    ]
    assert {cast(str, zeile[0].value) for zeile in statistikwerte} == {
        "System-IAT nach Gl. 3.16",
        "Zwischenankunftszeit",
        "Bearbeitungszeit",
        "Vereinfachte Start-zu-Start-Zeitspanne",
        "Wartezeit",
    }
    assert all(isinstance(zeile[2].value, int) for zeile in statistikwerte)
    assert all(isinstance(zeile[3].value, Real) for zeile in statistikwerte)
    assert all(zeile[3].number_format == "0.##" for zeile in statistikwerte)
    assert any(
        zelle.value == "Transformationshistorie D → aktiver Datensatz T"
        for zeile in datenblatt.iter_rows()
        for zelle in zeile
    )
    analysewerte = {
        str(zelle.value)
        for zeile in arbeitsmappe["Analyseergebnisse"].iter_rows()
        for zelle in zeile
        if zelle.value is not None
    }
    assert "Token-Based Replay · Gleichung 3.14" in analysewerte
    assert "99,75 %" in analysewerte
    assert "0,8" in analysewerte
    assert "Soll-/Ist-Abweichungen" in analysewerte
    assert "Ergänzende Performance · Busy Ratio" in analysewerte
    assert any(
        zelle.value == "Vereinfachte Start-zu-Start-Zeitspannen"
        for zeile in arbeitsmappe["Annahmen & offene Punkte"].iter_rows()
        for zelle in zeile
    )
    assert (
        len(
            {
                cast(float, datenblatt.column_dimensions[spalte].width)
                for spalte in ("A", "B", "C", "D", "E", "F", "G")
            }
        )
        > 2
    )


def test_xlsx_renderer_bettet_prozessgrafik_als_png_ein(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    report = build_report_data(_k_stern())
    projekt_id = report["projekt"]["projekt_id"]
    analyse_id = report["prozessdarstellung"]["process_mining_analyse_id"]
    ordner = tmp_path / "projects" / projekt_id / "process_mining"
    ordner.mkdir(parents=True)
    (ordner / f"{analyse_id}.model.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"></svg>',
        encoding="utf-8",
    )
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\rIDAT\x08\xd7c\xf8\xcf\xc0\xf0\x1f\x00"
        b"\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    monkeypatch.setattr(xlsx_modul, "_svg_zu_png", lambda _: png)

    inhalt = render_report_xlsx(resolve_report_assets(report, workspace_root=tmp_path))
    arbeitsmappe = load_workbook(BytesIO(inhalt))

    assert len(arbeitsmappe["Prozessmodell"]._images) == 1  # pyright: ignore[reportAttributeAccessIssue]


def test_aelteres_k_stern_ohne_optionale_felder_bleibt_renderbar(tmp_path: Path) -> None:
    report = resolve_report_assets(
        build_report_data(_k_stern(neue_felder=False)), workspace_root=tmp_path
    )

    assert report["warteschlangen"]["wartestellenhinweise"] == []
    assert report["ressourcen"]["aktivitaet_ressourcen"] == []
    assert report["ressourcen"]["manuelle_aktivitaet_ressourcen"] == []
    html = render_report_html(report)
    assert "<!DOCTYPE html>" in html
    assert "Conformance Checking · Gleichung 3.14" not in html
    assert "Soll-/Ist-Abweichungen" not in html
    assert "Ergänzende Performance · Busy Ratio" not in html
    assert "Vereinfachte Start-zu-Start-Zeitspannen" not in html
    ziel = render_report_pdf(report, tmp_path / "alt.pdf")
    assert ziel.read_bytes().startswith(b"%PDF-")


def test_automatische_und_manuelle_ressourcen_werden_gleichwertig_mit_ursprung_berichtet() -> None:
    manuell = build_report_data(_k_stern())
    automatisch_k = _k_stern()
    bestandteile = cast(list[dict[str, Any]], automatisch_k["modellbestandteile"])
    ressourcen = next(wert for wert in bestandteile if wert["bestandteil_id"] == "ressourcen")
    information = ressourcen["urspruenglicher_bestandteil"]["informationen"][0]
    information["wert"]["modus"] = "automatisch"
    information["wert"]["herkunft"] = "kanonische Ressourcenspalte in E*"

    automatisch = build_report_data(automatisch_k)

    assert (
        automatisch["ressourcen"]["aktivitaet_ressourcen"]
        == manuell["ressourcen"]["aktivitaet_ressourcen"]
    )
    assert automatisch["ressourcen"]["zuordnungsmodus"] == "automatisch"
    assert automatisch["ressourcen"]["zuordnungsherkunft"].startswith("kanonische")


def test_ressourcenanzeige_begrenzt_50_werte_ohne_die_a_g_projektion_zu_verkuerzen(
    tmp_path: Path,
) -> None:
    k_stern = _k_stern()
    ressourcennamen = [
        f"Kapazitaet-{index:02d}-mit-sehr-langer-eindeutiger-Bezeichnung-{'x' * 48}"
        for index in range(1, 51)
    ]
    bestandteile = cast(list[dict[str, Any]], k_stern["modellbestandteile"])
    ressourcen = next(wert for wert in bestandteile if wert["bestandteil_id"] == "ressourcen")
    information = ressourcen["urspruenglicher_bestandteil"]["informationen"][0]
    information["wert"]["zuordnungen"] = [
        {
            "aktivitaet": "Aktivitaet-mit-sehr-langer-Bezeichnung",
            "ressourcen": ressourcennamen,
        }
    ]
    vorher = copy.deepcopy(k_stern)

    report = build_report_data(k_stern)

    assert k_stern == vorher
    assert report["ressourcen"]["aktivitaet_ressourcen"][0]["ressourcen"] == ressourcennamen
    anzeige = report["ressourcen"]["aktivitaet_ressourcen_anzeige"][0]
    assert anzeige["ressourcen"] == ressourcennamen[:10]
    assert anzeige["weitere_ressourcen"] == 40
    html = render_report_html(report)
    zuordnung_html = html.split("<h2>Aktivität-Ressourcen-Zuordnungen</h2>", 1)[1].split(
        "</table>", 1
    )[0]
    assert ressourcennamen[9] in zuordnung_html
    assert ressourcennamen[10] not in zuordnung_html
    assert "+40 weitere" in zuordnung_html

    ziel = render_report_pdf(report, tmp_path / "ressourcen-50.pdf")
    css = (template_verzeichnis() / "report_pdf.css").read_text(encoding="utf-8")
    assert ziel.read_bytes().startswith(b"%PDF-")
    assert ".pdf-resource-names" in css
    assert "overflow-wrap: anywhere" in css
    assert "min-width: 0" in css


def test_build_report_data_weist_unvollstaendige_struktur_kontrolliert_ab() -> None:
    k_stern = _k_stern()
    k_stern["modellbestandteile"] = []
    with pytest.raises(ReportDataFehler, match="Fehlend"):
        build_report_data(k_stern)


def test_asset_resolver_loest_realen_workspace_und_optionale_svgs_auf(tmp_path: Path) -> None:
    report = build_report_data(_k_stern())
    projekt_id = report["projekt"]["projekt_id"]
    analyse_id = report["prozessdarstellung"]["process_mining_analyse_id"]
    ordner = tmp_path / "projects" / projekt_id / "process_mining"
    ordner.mkdir(parents=True)
    (ordner / f"{analyse_id}.model.svg").write_text(
        '<?xml version="1.0"?><svg><text>Modell</text></svg>', encoding="utf-8"
    )
    (ordner / f"{analyse_id}.dfg.svg").write_text("<svg><text>DFG</text></svg>", encoding="utf-8")

    aufgeloest = resolve_report_assets(report, workspace_root=tmp_path)

    assert aufgeloest["prozessdarstellung"]["svg_inline"].startswith("<svg")
    assert aufgeloest["prozessdarstellung"]["dfg_svg_inline"].startswith("<svg")
    assert aufgeloest["prozessdarstellung"]["process_tree_svg_inline"] is None
    assert aufgeloest["prozessdarstellung"]["assets"] == {
        "modell_svg": True,
        "dfg_svg": True,
        "process_tree_svg": False,
    }


def test_asset_resolver_toleriert_fehlende_und_verwirft_ungueltige_svgs(
    tmp_path: Path,
) -> None:
    report = build_report_data(_k_stern())
    ohne = resolve_report_assets(report, workspace_root=tmp_path)
    assert not any(ohne["prozessdarstellung"]["assets"].values())
    projekt_id = report["projekt"]["projekt_id"]
    analyse_id = report["prozessdarstellung"]["process_mining_analyse_id"]
    ordner = tmp_path / "projects" / projekt_id / "process_mining"
    ordner.mkdir(parents=True)
    (ordner / f"{analyse_id}.model.svg").write_text("kein SVG", encoding="utf-8")

    with pytest.raises(ReportAssetFehler, match="kein gültiges SVG"):
        resolve_report_assets(report, workspace_root=tmp_path)


def test_html_renderer_bettet_die_einzige_css_quelle_und_svgs_ein(tmp_path: Path) -> None:
    report = resolve_report_assets(build_report_data(_k_stern()), workspace_root=tmp_path)
    css = (template_verzeichnis() / "report_html.css").read_text(encoding="utf-8")

    html = render_report_html(report)

    assert '<link rel="stylesheet" href="report_html.css">' not in html
    assert f"<style>\n{css}\n</style>" in html
    assert "Übergangswartezeiten aus Schritt 7" in html
    assert "Aktivität-Ressourcen-Zuordnungen" in html
    assert "menschlich bestätigte Zuordnung in Schritt 7" in html
    assert "Zwischenankunftszeit" in html
    assert "Ist-Ende(A) − Ist-Start(A)" in html
    assert "Für spätere manuelle Berechnung vorgesehen" in html
    assert "Conformance Checking" in html
    assert "Produzierte Tokens pT" in html
    assert "dT · Fertigstellungsabweichung" in html
    assert "mittelwert sekunden" in html
    assert "20.0" not in html
    assert "Ergänzende Performance · Busy Ratio" in html
    assert "<td>0,8</td>" in html
    assert "Vereinfachte Start-zu-Start-Zeitspannen" in html
    assert "Potenzielle Wartestellen · Gleichung 3.15" in html


def test_pdf_renderer_verwendet_pdf_template_und_css(tmp_path: Path) -> None:
    report = resolve_report_assets(build_report_data(_k_stern()), workspace_root=tmp_path)
    template = template_verzeichnis() / "report_pdf.html"
    css = template_verzeichnis() / "report_pdf.css"
    assert 'href="report_pdf.css"' in template.read_text(encoding="utf-8")
    assert "@page" in css.read_text(encoding="utf-8")

    ziel = render_report_pdf(report, tmp_path / "bericht.pdf")

    assert ziel.read_bytes().startswith(b"%PDF-")


def test_service_uebergibt_identische_gemeinsame_reportdaten_an_alle_renderer(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    k_stern = _k_stern()
    aufgeloest = {"report_data_version": 1, "gemeinsam": object()}
    aufrufe = {"build": 0, "resolve": 0}
    renderer_ids: list[int] = []

    def build(wert, **metadaten):  # type: ignore[no-untyped-def]
        assert wert is k_stern
        assert metadaten == {
            "projektbezeichnung": "Fördertechnik Süd / ÄÖÜ",
            "softwareversion": "1.2",
        }
        aufrufe["build"] += 1
        return {"report_data_version": 1}

    def resolve(wert, *, workspace_root):  # type: ignore[no-untyped-def]
        assert wert == {"report_data_version": 1}
        assert workspace_root == tmp_path
        aufrufe["resolve"] += 1
        return aufgeloest

    def html_renderer(wert):  # type: ignore[no-untyped-def]
        renderer_ids.append(id(wert))
        return "<html></html>"

    def pdf_renderer(wert, ziel):  # type: ignore[no-untyped-def]
        renderer_ids.append(id(wert))
        Path(ziel).write_bytes(b"%PDF-test")
        return Path(ziel)

    def xlsx_renderer(wert):  # type: ignore[no-untyped-def]
        renderer_ids.append(id(wert))
        return b"PK-xlsx-test"

    monkeypatch.setattr(ausgabe_modul, "build_report_data", build)
    monkeypatch.setattr(ausgabe_modul, "resolve_report_assets", resolve)
    monkeypatch.setattr(ausgabe_modul, "render_report_html", html_renderer)
    monkeypatch.setattr(ausgabe_modul, "render_report_pdf", pdf_renderer)
    monkeypatch.setattr(ausgabe_modul, "render_report_xlsx", xlsx_renderer)
    validierungen = SimpleNamespace(uebergabe_schritt10=lambda *_: k_stern)
    projekte = SimpleNamespace(
        projekt_laden=lambda projekt_id: SimpleNamespace(
            projekt_id=projekt_id, bezeichnung="Fördertechnik Süd / ÄÖÜ"
        )
    )
    service = ModellausgabeService(
        cast(ModellvalidierungService, validierungen),
        cast(Any, projekte),
        WorkspaceKonfiguration(tmp_path),
    )

    ergebnis = service.erzeugen(
        validierungslauf_id=UUID(str(k_stern["validierungslauf_id"])),
        projekt_id=UUID(str(k_stern["projekt_id"])),
        k_stern_id=UUID(str(k_stern["k_stern_id"])),
        html=True,
        pdf=True,
        xlsx=True,
    )

    assert aufrufe == {"build": 1, "resolve": 1}
    assert renderer_ids == [id(aufgeloest), id(aufgeloest), id(aufgeloest)]
    assert ergebnis.report_html == b"<html></html>"
    assert ergebnis.report_pdf == b"%PDF-test"
    assert ergebnis.report_xlsx == b"PK-xlsx-test"
    assert ergebnis.pdf_dateiname == "Konzeptionelles Modell Fördertechnik Süd ÄÖÜ.pdf"
    assert ergebnis.xlsx_dateiname == "Konzeptionelles Modell Fördertechnik Süd ÄÖÜ.xlsx"


def test_service_uebersetzt_reportingfehler_in_die_anwendungsschicht(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    k_stern = _k_stern()
    monkeypatch.setattr(
        ausgabe_modul,
        "build_report_data",
        lambda _wert, **_metadaten: (_ for _ in ()).throw(ReportDataFehler("inkompatibel")),
    )
    service = ModellausgabeService(
        cast(
            ModellvalidierungService,
            SimpleNamespace(uebergabe_schritt10=lambda *_: k_stern),
        ),
        cast(
            Any,
            SimpleNamespace(
                projekt_laden=lambda projekt_id: SimpleNamespace(
                    projekt_id=projekt_id, bezeichnung="Reporting"
                )
            ),
        ),
        WorkspaceKonfiguration(tmp_path),
    )

    with pytest.raises(Importintegritaetsfehler, match="inkompatibel"):
        service.erzeugen(
            validierungslauf_id=UUID(str(k_stern["validierungslauf_id"])),
            projekt_id=UUID(str(k_stern["projekt_id"])),
            k_stern_id=UUID(str(k_stern["k_stern_id"])),
            html=True,
            pdf=False,
        )


def test_service_verwendet_keinen_projektnamen_einer_fremden_id(tmp_path: Path) -> None:
    k_stern = _k_stern()
    projekt_id = UUID(str(k_stern["projekt_id"]))
    service = ModellausgabeService(
        cast(
            ModellvalidierungService,
            SimpleNamespace(uebergabe_schritt10=lambda *_: k_stern),
        ),
        cast(
            Any,
            SimpleNamespace(
                projekt_laden=lambda _: SimpleNamespace(
                    projekt_id=uuid4(), bezeichnung="Fremdes Projekt"
                )
            ),
        ),
        WorkspaceKonfiguration(tmp_path),
    )

    with pytest.raises(Importintegritaetsfehler, match="Projektbezeichnung"):
        service.erzeugen(
            validierungslauf_id=UUID(str(k_stern["validierungslauf_id"])),
            projekt_id=projekt_id,
            k_stern_id=UUID(str(k_stern["k_stern_id"])),
            html=True,
            pdf=False,
        )


def test_alte_pdf_excel_und_browser_print_logik_ist_nicht_mehr_produktiv() -> None:
    assert not hasattr(ausgabe_modul, "_pdf_erzeugen")
    assert not hasattr(ausgabe_modul, "_excel_erzeugen")
    ui_quelltext = Path("src/framework_mvp/ui/pages/modellausgabe.py").read_text(encoding="utf-8")
    assert "window.print" not in ui_quelltext
    assert "excel_xlsx" not in ui_quelltext
