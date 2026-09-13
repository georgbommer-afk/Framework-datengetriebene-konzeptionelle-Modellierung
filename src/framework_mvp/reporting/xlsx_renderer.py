"""Professionelle XLSX-Ausgabe aus den gemeinsamen Reportdaten von Schritt 10."""

from __future__ import annotations

import logging
import math
import shutil
import subprocess
from collections.abc import Iterable, Mapping, Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from framework_mvp.formatierung import formatiere_fachwert
from framework_mvp.reporting.report_data import REPORT_DATA_VERSION

LOGGER = logging.getLogger(__name__)

SHEET_NAMES = (
    "Übersicht",
    "Problemstellung & Systemgrenze",
    "Zielsetzung & KPIs",
    "Prozessmodell",
    "Systemelemente & Merkmale",
    "Annahmen & offene Punkte",
    "Daten & Datenauswahl",
    "Analyseergebnisse",
    "Validierung",
    "Nachvollziehbarkeit",
)

NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT_BLUE = "EEF5FA"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "A6A6A6"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
THIN_GRAY = Side(style="thin", color="D9E1F2")
FONT_FAMILY = "Calibri"


class XlsxRenderingFehler(RuntimeError):
    """Kennzeichnet eine nicht erzeugbare XLSX-Reportdatei."""


def _mapping(wert: Any) -> Mapping[str, Any]:
    return wert if isinstance(wert, Mapping) else {}


def _liste(wert: Any) -> list[Any]:
    if wert is None or wert == "":
        return []
    if isinstance(wert, list):
        return wert
    if isinstance(wert, tuple):
        return list(wert)
    return [wert]


def _label(name: Any) -> str:
    return str(name).replace("_", " ").strip().capitalize()


def _anzeige(wert: Any) -> str:
    """Formatiert strukturierte Werte lesbar, ohne JSON-Rohdump."""
    if wert is None or wert == "":
        return "Nicht im finalen Modell ausgewiesen"
    if isinstance(wert, bool):
        return "Ja" if wert else "Nein"
    if isinstance(wert, Mapping):
        if not wert:
            return "Nicht im finalen Modell ausgewiesen"
        return "\n".join(f"{_label(name)}: {_anzeige(inhalt)}" for name, inhalt in wert.items())
    if isinstance(wert, (list, tuple)):
        if not wert:
            return "Nicht im finalen Modell ausgewiesen"
        return "\n".join(f"• {_anzeige(inhalt)}" for inhalt in wert)
    return str(formatiere_fachwert(wert))


def _kurz(wert: Any, laenge: int = 240) -> str:
    text = _anzeige(wert).replace("\n", " ")
    return text if len(text) <= laenge else f"{text[: laenge - 1].rstrip()}…"


def _sheet_start(ws: Worksheet, titel: str, untertitel: str) -> int:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:H1")
    ws["A1"] = titel
    ws["A1"].font = Font(name=FONT_FAMILY, size=20, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:H2")
    ws["A2"] = untertitel
    ws["A2"].font = Font(name=FONT_FAMILY, size=12, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 26
    for spalte in range(1, 9):
        ws.cell(3, spalte).border = Border(bottom=Side(style="medium", color=NAVY))
    if ws.sheet_properties.pageSetUpPr is not None:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if ws.sheet_properties.outlinePr is not None:
        ws.sheet_properties.outlinePr.summaryBelow = True
    return 5


def _abschnitt(ws: Worksheet, zeile: int, titel: str, breite: int = 8) -> int:
    ws.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=breite)
    zelle = ws.cell(zeile, 1, titel)
    zelle.fill = PatternFill("solid", fgColor=BLUE)
    zelle.font = Font(name=FONT_FAMILY, size=14, bold=True, color=NAVY)
    zelle.alignment = Alignment(vertical="center")
    ws.row_dimensions[zeile].height = 22
    return zeile + 1


def _paare(ws: Worksheet, zeile: int, paare: Iterable[tuple[str, Any]]) -> int:
    for bezeichnung, wert in paare:
        ws.cell(zeile, 1, bezeichnung)
        ws.cell(zeile, 1).font = Font(name=FONT_FAMILY, size=12, bold=True, color=NAVY)
        ws.merge_cells(start_row=zeile, start_column=2, end_row=zeile, end_column=8)
        ws.cell(zeile, 2, _anzeige(wert))
        for spalte in range(1, 9):
            zelle = ws.cell(zeile, spalte)
            zelle.alignment = Alignment(wrap_text=True, vertical="top")
            zelle.border = Border(bottom=THIN_GRAY)
        ws.row_dimensions[zeile].height = max(20, min(90, 15 * (_anzeige(wert).count("\n") + 1)))
        zeile += 1
    return zeile


def _status_fill(wert: Any) -> PatternFill | None:
    text = str(wert).casefold()
    if any(teil in text for teil in ("validiert", "berechnet", "übernommen", "bestaetigt")):
        return PatternFill("solid", fgColor=GREEN)
    if any(teil in text for teil in ("offen", "unsicher", "nicht berechenbar")):
        return PatternFill("solid", fgColor=YELLOW)
    if any(teil in text for teil in ("fehler", "ungültig", "abgelehnt")):
        return PatternFill("solid", fgColor=RED)
    return None


def _tabelle(
    ws: Worksheet,
    zeile: int,
    spalten: Sequence[tuple[str, str]],
    daten: Iterable[Mapping[str, Any]],
    *,
    autofilter: bool = False,
    zahlenformate: Mapping[str, str] | None = None,
) -> int:
    datenzeilen = list(daten)
    zahlenformate = zahlenformate or {}
    for index, (titel, _) in enumerate(spalten, 1):
        zelle = ws.cell(zeile, index, titel)
        zelle.fill = PatternFill("solid", fgColor=NAVY)
        zelle.font = Font(name=FONT_FAMILY, size=12, bold=True, color=WHITE)
        zelle.alignment = Alignment(wrap_text=True, vertical="center")
    kopfzeile = zeile
    zeile += 1
    if not datenzeilen:
        ws.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=len(spalten))
        ws.cell(zeile, 1, "Für dieses Projekt sind im finalen Modell keine Einträge vorhanden.")
        ws.cell(zeile, 1).font = Font(name=FONT_FAMILY, size=12, italic=True, color="666666")
        zeile += 1
    else:
        for datensatz in datenzeilen:
            maximale_zeilen = 1
            for index, (_, schluessel) in enumerate(spalten, 1):
                rohwert = datensatz.get(schluessel)
                text = _anzeige(rohwert)
                zellenwert = (
                    rohwert
                    if schluessel in zahlenformate
                    and isinstance(rohwert, (int, float))
                    and not isinstance(rohwert, bool)
                    else text
                )
                zelle = ws.cell(zeile, index, zellenwert)
                zelle.font = Font(name=FONT_FAMILY, size=12)
                if schluessel in zahlenformate and isinstance(zellenwert, (int, float)):
                    zelle.number_format = zahlenformate[schluessel]
                zelle.alignment = Alignment(wrap_text=True, vertical="top")
                zelle.border = Border(bottom=THIN_GRAY)
                zeichen_pro_zeile = 24 if index == 1 else 30
                benoetigte_zeilen = sum(
                    max(1, math.ceil(len(abschnitt) / zeichen_pro_zeile))
                    for abschnitt in text.splitlines()
                )
                maximale_zeilen = max(maximale_zeilen, benoetigte_zeilen)
                if index == len(spalten):
                    fuellung = _status_fill(zelle.value)
                    if fuellung is not None:
                        zelle.fill = fuellung
            ws.row_dimensions[zeile].height = min(405, max(24, 14 * maximale_zeilen + 4))
            zeile += 1
    if autofilter and datenzeilen:
        ws.auto_filter.ref = f"A{kopfzeile}:{get_column_letter(len(spalten))}{zeile - 1}"
    return zeile


def _informationen(report: Mapping[str, Any], ids: set[str]) -> list[dict[str, Any]]:
    ergebnis: list[dict[str, Any]] = []
    for bestandteil in _liste(report.get("modellbestandteile")):
        if not isinstance(bestandteil, Mapping) or bestandteil.get("bestandteil_id") not in ids:
            continue
        for information in _liste(bestandteil.get("informationen")):
            if isinstance(information, Mapping):
                ergebnis.append(
                    {
                        "bestandteil": bestandteil.get("bezeichnung"),
                        "referenz": information.get("strukturreferenz"),
                        "wert": information.get("wert"),
                        "quelle": information.get("herkunftsartefakt"),
                        "status": bestandteil.get("validierungsstatus_anzeige"),
                    }
                )
    return ergebnis


def _svg_zu_png(svg: str) -> bytes | None:
    """Konvertiert SVG bevorzugt in-process, optional über vorhandenes rsvg-convert."""
    try:
        from cairosvg import svg2png  # type: ignore[import-not-found]

        return bytes(svg2png(bytestring=svg.encode("utf-8"), output_width=1800))
    except (ImportError, OSError, ValueError):
        programm = shutil.which("rsvg-convert")
        if programm is None:
            return None
        try:
            prozess = subprocess.run(
                [programm, "--format=png", "--width=1800"],
                input=svg.encode("utf-8"),
                capture_output=True,
                check=True,
                timeout=20,
            )
        except (OSError, subprocess.SubprocessError):
            return None
        return prozess.stdout or None


def _prozessbild(ws: Worksheet, zeile: int, report: Mapping[str, Any]) -> int:
    prozess = _mapping(report.get("prozessdarstellung"))
    svg = next(
        (
            wert
            for wert in (
                prozess.get("svg_inline"),
                prozess.get("dfg_svg_inline"),
                prozess.get("process_tree_svg_inline"),
            )
            if isinstance(wert, str) and wert.strip()
        ),
        None,
    )
    if svg is None:
        return _paare(ws, zeile, [("Visualisierung", "Keine Prozessgrafik im K*-Asset verfügbar.")])
    png = _svg_zu_png(svg)
    if png is None:
        LOGGER.warning(
            "Die Prozessgrafik konnte für die XLSX-Ausgabe nicht in PNG umgewandelt werden."
        )
        return _paare(
            ws,
            zeile,
            [("Visualisierung", "Die Prozessgrafik konnte nicht in Excel eingebettet werden.")],
        )
    try:
        bild = Image(BytesIO(png))
        verfuegbare_breite = (
            sum(
                float(ws.column_dimensions[get_column_letter(index)].width or 13.0) * 7.0
                for index in range(1, 9)
            )
            - 24.0
        )
        faktor = min(1.0, verfuegbare_breite / bild.width, 720 / bild.height)
        bild.width = int(bild.width * faktor)
        bild.height = int(bild.height * faktor)
        ws.add_image(bild, f"A{zeile}")
        belegte_zeilen = max(10, int(bild.height / 20) + 2)
        for index in range(zeile, zeile + belegte_zeilen):
            ws.row_dimensions[index].height = 20
        return zeile + belegte_zeilen
    except (OSError, ValueError) as fehler:
        LOGGER.warning("Die PNG-Prozessgrafik konnte nicht eingebettet werden: %s", fehler)
        return _paare(
            ws,
            zeile,
            [("Visualisierung", "Die Prozessgrafik konnte nicht in Excel eingebettet werden.")],
        )


def _uebersicht(ws: Worksheet, report: Mapping[str, Any]) -> None:
    projekt = _mapping(report.get("projekt"))
    modell = _mapping(report.get("modell"))
    umfang = _mapping(report.get("modellumfang"))
    ziel = _mapping(report.get("zielsetzung"))
    validierung = _mapping(report.get("validierung"))
    dokument = _mapping(report.get("dokument"))
    zeile = _sheet_start(ws, SHEET_NAMES[0], "Management Summary des finalen K*")
    zeile = _abschnitt(ws, zeile, "Projekt und Modell")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Projekt", projekt.get("bezeichnung") or dokument.get("titel")),
            ("Systemtyp", umfang.get("systemtyp_anzeige") or umfang.get("systemtyp")),
            (
                "Systemgegenstand",
                umfang.get("bereich_aus_systemprofil")
                or _mapping(umfang.get("systemklassifikation")).get("bereich"),
            ),
            ("Problemstellung", _kurz(_mapping(report.get("problemstellung")).get("text"), 400)),
            ("Zielsetzung", ziel.get("individuelles_ziel") or ziel.get("untersuchungszwecke")),
            ("Validierungsstatus", validierung.get("status_anzeige")),
            ("Fachliche Validierung", validierung.get("validierungsvermerk")),
            ("Modellstand", modell.get("erstellt_am")),
            ("Framework-Version", dokument.get("softwareversion")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Wichtigste Modellbestandteile")
    ressourcen = _mapping(report.get("ressourcen"))
    zusammenfassung = [
        {
            "kategorie": "Entitäten/Objekte",
            "inhalt": _mapping(report.get("entitaeten")).get("objekte_gueter"),
        },
        {
            "kategorie": "Ressourcen",
            "inhalt": ressourcen.get("event_log_ressourcen") or ressourcen.get("systemressourcen"),
        },
        {
            "kategorie": "Aktivitäten",
            "inhalt": _mapping(report.get("aktivitaeten")).get("sichtbare_aktivitaeten"),
        },
        {
            "kategorie": "Warteschlangen",
            "inhalt": [
                wert.get("bezeichnung")
                or f"{wert.get('vorgaengeraktivitaet', '')} → {wert.get('folgeaktivitaet', '')}"
                for wert in _liste(
                    _mapping(report.get("warteschlangen")).get("bestaetigte_warteschlangen")
                )
                if isinstance(wert, Mapping)
            ],
        },
        {
            "kategorie": "Zustände",
            "inhalt": [
                wert.get("wert")
                for wert in _informationen(report, {"entitaeten", "aktivitaeten"})
                if "zustand" in str(wert.get("referenz", "")).casefold()
            ],
        },
        {
            "kategorie": "Fallidentifikation",
            "inhalt": _mapping(report.get("entitaeten")).get("fallidentifikation"),
        },
        {
            "kategorie": "Prozesssteuerung/Logik",
            "inhalt": _mapping(report.get("annahmen")).get("modellierungsentscheidungen"),
        },
        {
            "kategorie": "Eingaben",
            "inhalt": [
                f"{wert.get('bezeichnung', 'Experimenteller Faktor')} "
                f"({wert.get('bezugstyp_anzeige', 'Bezug')} "
                f"{wert.get('konkreter_bezug', '')}): "
                f"{wert.get('wertebereich_anzeige', '')}"
                for wert in _liste(_mapping(report.get("eingaben")).get("experimentelle_faktoren"))
                if isinstance(wert, Mapping)
            ],
        },
        {
            "kategorie": "Ausgaben",
            "inhalt": [
                wert.get("bezeichnung")
                for wert in _liste(_mapping(report.get("ausgaben")).get("ausgewaehlte_kpis"))
                if isinstance(wert, Mapping)
            ],
        },
        {
            "kategorie": "Annahmen",
            "inhalt": _mapping(report.get("annahmen")).get("schwellwert_auswirkung"),
        },
    ]
    zeile = _tabelle(
        ws,
        zeile,
        (("Kategorie", "kategorie"), ("Finaler Inhalt", "inhalt")),
        zusammenfassung,
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Navigation")
    for ziel in SHEET_NAMES[1:]:
        ws.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=3)
        zelle = ws.cell(zeile, 1, ziel)
        zelle.hyperlink = f"#'{ziel}'!A1"
        zelle.style = "Hyperlink"
        zeile += 1
    ws.column_dimensions["A"].width = 28
    for spalte in range(2, 9):
        ws.column_dimensions[get_column_letter(spalte)].width = 16


def _problem_und_grenze(ws: Worksheet, report: Mapping[str, Any]) -> None:
    problem = _mapping(report.get("problemstellung"))
    umfang = _mapping(report.get("modellumfang"))
    zeile = _sheet_start(ws, SHEET_NAMES[1], "Fachlicher Auftrag und Abgrenzung")
    zeile = _abschnitt(ws, zeile, "Problemstellung")
    zeile = _paare(ws, zeile, (("Problem", problem.get("text")),))
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Systemgrenze und Umfang")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Systemgrenze", umfang.get("systemgrenze")),
            ("Modellumfang", umfang.get("sichtbare_aktivitaeten")),
            ("Detaillierungsgrad", umfang.get("detaillierungsgrad")),
            ("Einbezogener Bereich", umfang.get("bereich_aus_systemprofil")),
            ("Systemklassifikation", umfang.get("systemklassifikation")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Start- und Endaktivitäten")
    _tabelle(
        ws,
        zeile,
        (("Typ", "typ"), ("Aktivität", "aktivitaet"), ("Häufigkeit", "haeufigkeit")),
        [
            {"typ": typ, **dict(eintrag)}
            for typ, schluessel in (("Start", "startaktivitaeten"), ("Ende", "endaktivitaeten"))
            for eintrag in _liste(umfang.get(schluessel))
            if isinstance(eintrag, Mapping)
        ],
        zahlenformate={"haeufigkeit": "#,##0"},
    )


def _ziele(ws: Worksheet, report: Mapping[str, Any]) -> None:
    ziel = _mapping(report.get("zielsetzung"))
    ausgaben = _mapping(report.get("ausgaben"))
    eingaben = _mapping(report.get("eingaben"))
    zeile = _sheet_start(ws, SHEET_NAMES[2], "Ziele, Zielgrößen und Kennzahlen")
    zeile = _abschnitt(ws, zeile, "Zielbild")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Untersuchungszwecke", ziel.get("untersuchungszwecke")),
            ("Individuelles Ziel", ziel.get("individuelles_ziel")),
            ("Logistische Zielgrößen", ziel.get("logistische_zielgroessen")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Ausgewählte KPIs und Ergebnisse")
    kpis = [
        dict(wert) for wert in _liste(ausgaben.get("kpi_ergebnisse")) if isinstance(wert, Mapping)
    ]
    vorhandene = {str(wert.get("kpi_id")) for wert in kpis}
    kpis.extend(
        {
            "kpi_id": wert.get("id"),
            "bezeichnung": wert.get("bezeichnung"),
            "status_anzeige": "Ausgewählt",
            "ergebnis": None,
            "einheit": None,
            "formel": None,
            "bezugsmenge": None,
        }
        for wert in _liste(ausgaben.get("ausgewaehlte_kpis"))
        if isinstance(wert, Mapping) and str(wert.get("id")) not in vorhandene
    )
    for kpi in kpis:
        if not isinstance(kpi.get("ergebnis"), (int, float)):
            kpi["ergebnis"] = kpi.get("ergebnis_anzeige")
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("Bezeichnung", "bezeichnung"),
            ("Ergebnis", "ergebnis"),
            ("Einheit", "einheit"),
            ("Bezugsmenge", "bezugsmenge"),
            ("Berechnungslogik", "formel"),
            ("Status", "status_anzeige"),
        ),
        kpis,
        autofilter=True,
        zahlenformate={"ergebnis": "0.##"},
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Experimentelle Faktoren")
    _tabelle(
        ws,
        zeile,
        (
            ("Bezugstyp", "bezugstyp_anzeige"),
            ("Konkreter Bezug", "konkreter_bezug"),
            ("Experimenteller Faktor", "bezeichnung"),
            ("Art", "art_anzeige"),
            ("Wertebereich / Ausprägungen", "wertebereich_anzeige"),
        ),
        [
            dict(wert)
            for wert in _liste(eingaben.get("experimentelle_faktoren"))
            if isinstance(wert, Mapping)
        ],
        autofilter=True,
    )


def _prozess(ws: Worksheet, report: Mapping[str, Any]) -> None:
    prozess = _mapping(report.get("prozessdarstellung"))
    annahmen = _mapping(report.get("annahmen"))
    zeile = _sheet_start(
        ws, SHEET_NAMES[3], "Eingebettete Visualisierung des finalen Prozessmodells"
    )
    zeile = _abschnitt(ws, zeile, "Modell und Mining-Parameter")
    zeile = _paare(
        ws,
        zeile,
        (
            (
                "Notation",
                prozess.get("notation_anzeige") or annahmen.get("prozessnotation_anzeige"),
            ),
            ("Modellierungsentscheidungen", annahmen.get("modellierungsentscheidungen")),
            ("Schwellwert-Auswirkung", annahmen.get("schwellwert_auswirkung")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Prozessvisualisierung")
    ws.column_dimensions["A"].width = 30
    for spalte in range(2, 9):
        ws.column_dimensions[get_column_letter(spalte)].width = 24
    _prozessbild(ws, zeile, report)
    ws.sheet_view.zoomScale = 80


def _systemelemente(ws: Worksheet, report: Mapping[str, Any]) -> None:
    zeile = _sheet_start(
        ws, SHEET_NAMES[4], "Finale Entitäten, Aktivitäten, Ressourcen und Merkmale"
    )
    zeile = _abschnitt(ws, zeile, "Strukturierte Modellinhalte")
    entitaeten = _mapping(report.get("entitaeten"))
    aktivitaeten = _mapping(report.get("aktivitaeten"))
    ressourcen = _mapping(report.get("ressourcen"))
    warteschlangen = _mapping(report.get("warteschlangen"))
    daten: list[dict[str, Any]] = []

    def ergaenzen(kategorie: str, werte: Any, *, ursprung: str) -> None:
        for wert in _liste(werte):
            daten.append(
                {
                    "kategorie": kategorie,
                    "element": _kurz(wert, 500),
                    "typ": kategorie,
                    "beschreibung": _kurz(wert, 500),
                    "ursprung": ursprung,
                    "eigenschaften": None,
                    "status": "Fachlich validiert",
                }
            )

    ergaenzen("Entität/Objekt", entitaeten.get("objekte_gueter"), ursprung="Systemprofil")
    ergaenzen(
        "Aktivität",
        aktivitaeten.get("sichtbare_aktivitaeten"),
        ursprung="Finales Prozessmodell",
    )
    ergaenzen(
        "Ressource",
        ressourcen.get("event_log_ressourcen") or ressourcen.get("systemressourcen"),
        ursprung=(
            "Automatisch aus dem Event Log abgeleitet"
            if ressourcen.get("zuordnungsmodus") == "automatisch"
            else "In Schritt 7 fachlich zugeordnet"
        ),
    )
    ergaenzen(
        "Warteschlange",
        warteschlangen.get("bestaetigte_warteschlangen"),
        ursprung="Fachliche Validierung",
    )
    if entitaeten.get("fallidentifikation"):
        daten.append(
            {
                "kategorie": "Entität/Objekt",
                "element": "Fallidentifikation",
                "typ": "Fachlicher Fallbezug",
                "beschreibung": entitaeten.get("fallidentifikation"),
                "ursprung": "Systemprofil",
                "eigenschaften": None,
                "status": "Fachlich validiert",
            }
        )
    for faktor in _liste(_mapping(report.get("eingaben")).get("experimentelle_faktoren")):
        mapping = _mapping(faktor)
        daten.append(
            {
                "kategorie": "Eingabe",
                "element": mapping.get("bezeichnung"),
                "typ": "Experimenteller Faktor",
                "beschreibung": mapping.get("wertebereich_anzeige"),
                "ursprung": "Fachliche Validierung",
                "eigenschaften": {
                    "Bezugstyp": mapping.get("bezugstyp_anzeige"),
                    "Konkreter Bezug": mapping.get("konkreter_bezug"),
                },
                "status": "Fachlich validiert",
            }
        )
    ausgaben = _mapping(report.get("ausgaben"))
    for kpi in _liste(ausgaben.get("kpi_ergebnisse")):
        mapping = _mapping(kpi)
        daten.append(
            {
                "kategorie": "Ausgabe/KPI",
                "element": mapping.get("bezeichnung") or "KPI-Ergebnis",
                "typ": "Kennzahl",
                "beschreibung": mapping.get("ergebnis_anzeige"),
                "ursprung": "Ergebnisaggregation",
                "eigenschaften": {
                    "Formel": mapping.get("formel"),
                    "Bezugsmenge": mapping.get("bezugsmenge"),
                },
                "status": mapping.get("status_anzeige"),
            }
        )
    _tabelle(
        ws,
        zeile,
        (
            ("Kategorie", "kategorie"),
            ("Element/Merkmal", "element"),
            ("Typ", "typ"),
            ("Beschreibung", "beschreibung"),
            ("Ursprung", "ursprung"),
            ("Eigenschaften", "eigenschaften"),
            ("Status", "status"),
        ),
        daten,
        autofilter=True,
    )


def _annahmen(ws: Worksheet, report: Mapping[str, Any]) -> None:
    annahmen = _mapping(report.get("annahmen"))
    vereinfachungen = _mapping(report.get("vereinfachungen"))
    zeile = _sheet_start(ws, SHEET_NAMES[5], "Annahmen, Vereinfachungen und fachliche Restpunkte")
    zeile = _abschnitt(ws, zeile, "Modellierungsannahmen")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Modellierungsentscheidungen", annahmen.get("modellierungsentscheidungen")),
            ("Prozessnotation", annahmen.get("prozessnotation_anzeige")),
            ("Schwellwert-Auswirkung", annahmen.get("schwellwert_auswirkung")),
            (
                "Vereinfachte Start-zu-Start-Zeitspannen",
                vereinfachungen.get("vereinfachte_zeitspannen"),
            ),
        ),
    )
    explizite_inhalte = [
        {
            "typ": "Annahme",
            "bezug": "Gesamtsystem",
            "inhalt": wert.get("annahme"),
            "hintergrund": wert.get("hintergrund"),
        }
        for wert in _liste(annahmen.get("explizite_annahmen"))
        if isinstance(wert, Mapping)
    ]
    explizite_inhalte.extend(
        {
            "typ": "Vereinfachung",
            "bezug": wert.get("konkreter_bezug") or "Gesamtsystem",
            "inhalt": wert.get("beschreibung"),
            "hintergrund": None,
        }
        for wert in _liste(vereinfachungen.get("explizite_vereinfachungen"))
        if isinstance(wert, Mapping)
    )
    if explizite_inhalte:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Explizite Annahmen und Vereinfachungen")
        zeile = _tabelle(
            ws,
            zeile,
            (
                ("Typ", "typ"),
                ("Bezug", "bezug"),
                ("Inhalt", "inhalt"),
                ("Hintergrund", "hintergrund"),
            ),
            explizite_inhalte,
            autofilter=True,
        )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Fachliche Entscheidungen, Anpassungen und offene Punkte")
    entscheidungen: list[dict[str, Any]] = []
    for bestandteil in _liste(report.get("modellbestandteile")):
        if not isinstance(bestandteil, Mapping):
            continue
        for typ, schluessel in (
            ("Entscheidung", "fachliche_entscheidungen"),
            ("Anpassung", "fachliche_anpassungen"),
        ):
            for eintrag in _liste(bestandteil.get(schluessel)):
                mapping = _mapping(eintrag)
                entscheidungen.append(
                    {
                        "bestandteil": bestandteil.get("bezeichnung"),
                        "typ": typ,
                        "inhalt": mapping.get("fachlicher_inhalt")
                        or mapping.get("entscheidung")
                        or mapping,
                        "begruendung": mapping.get("begruendung"),
                        "status": bestandteil.get("validierungsstatus_anzeige"),
                    }
                )
    _tabelle(
        ws,
        zeile,
        (
            ("Bestandteil", "bestandteil"),
            ("Typ", "typ"),
            ("Inhalt", "inhalt"),
            ("Begründung", "begruendung"),
            ("Status", "status"),
        ),
        entscheidungen,
        autofilter=True,
    )


def _daten(ws: Worksheet, report: Mapping[str, Any]) -> None:
    daten = _mapping(report.get("daten"))
    zeile = _sheet_start(ws, SHEET_NAMES[6], "Datenbasis und zeitbezogene Datenauswahl")
    zeile = _abschnitt(ws, zeile, "Datenquellen")
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("Datenquelle", "bezeichnung"),
            ("Quellsystem", "quellsystem_anzeige"),
            ("Format", "format_anzeige"),
            ("Verwendung", "verwendung"),
        ),
        [dict(wert) for wert in _liste(daten.get("datenquellen")) if isinstance(wert, Mapping)],
        autofilter=True,
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Datenartefakte")
    event_log = _mapping(daten.get("event_log"))
    zeile = _paare(
        ws,
        zeile,
        (
            (
                "Zwischendatensatz",
                {
                    "Zeilen": _mapping(daten.get("zwischendatensatz")).get("zeilenanzahl"),
                    "Spalten": _mapping(daten.get("zwischendatensatz")).get("spaltenanzahl"),
                },
            ),
            (
                "Event Log",
                {
                    "Ereignisse": event_log.get("ereignisanzahl"),
                    "Fälle": event_log.get("fallanzahl"),
                    "Aktivitäten": event_log.get("aktivitaetsanzahl"),
                },
            ),
            ("Von", event_log.get("zeitraum_von_anzeige", event_log.get("zeitraum_von"))),
            ("Bis", event_log.get("zeitraum_bis_anzeige", event_log.get("zeitraum_bis"))),
        ),
    )
    aufbereitung = _mapping(daten.get("datenaufbereitung"))
    transformationshistorie = [
        dict(wert)
        for wert in _liste(aufbereitung.get("transformationshistorie"))
        if isinstance(wert, Mapping)
    ]
    if aufbereitung:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Transformationshistorie D → aktiver Datensatz T")
        zeile = _paare(
            ws,
            zeile,
            (("Fachliche Bedeutung", aufbereitung.get("fachliche_bedeutung")),),
        )
        zeile = _tabelle(
            ws,
            zeile,
            (
                ("Reihenfolge", "reihenfolge"),
                ("Transformationsart", "transformationsart"),
                ("Spalten", "betroffene_spalten"),
                ("Regel", "regel"),
                ("Ersatz-/Abstraktionswert", "ersatz_oder_abstraktionswert"),
                ("Wirkung", "wirkung"),
                ("Ergebnis", "ergebnis"),
            ),
            transformationshistorie,
        )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Datenprofile")
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("Datenquelle", "anzeigebezeichnung"),
            ("Zeilen", "zeilenanzahl"),
            ("Spalten", "spaltenanzahl"),
            ("Fehlwerte", "echte_fehlwerte"),
            ("Platzhalter", "textuelle_platzhalter"),
            ("Duplikate", "exakte_duplikate"),
        ),
        [dict(wert) for wert in _liste(daten.get("profile")) if isinstance(wert, Mapping)],
        zahlenformate={
            "zeilenanzahl": "#,##0",
            "spaltenanzahl": "#,##0",
            "echte_fehlwerte": "#,##0",
            "textuelle_platzhalter": "#,##0",
            "exakte_duplikate": "#,##0",
        },
    )
    zeile += 1
    zeitwahl = _mapping(daten.get("zeitbezogene_datenauswahl"))
    zeile = _abschnitt(ws, zeile, "Zwischenankunftszeit System · Gl. 3.16")
    system_iat = _mapping(zeitwahl.get("system_zwischenankunftszeit"))
    system_stats = _mapping(system_iat.get("statistik"))
    zeile = _paare(
        ws,
        zeile,
        (
            ("Systemeintritt", "Frühester Ist-Start je Fall"),
            ("Entitäten", system_iat.get("anzahl_entitaeten")),
            ("Anzahl Zwischenankunftszeiten", system_stats.get("anzahl")),
            ("Mittelwert (s)", system_stats.get("mittelwert_sekunden")),
            ("Median (s)", system_stats.get("median_sekunden")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Bearbeitungszeit · Gl. 3.3")
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("Aktivität", "aktivitaet"),
            ("Ressource", "ressource"),
            ("n", "anzahl"),
            ("Mittelwert (s)", "mittelwert_sekunden"),
            ("Median (s)", "median_sekunden"),
        ),
        [
            dict(wert)
            for wert in _liste(zeitwahl.get("bearbeitungszeiten_tabelle"))
            if isinstance(wert, Mapping)
        ],
        autofilter=True,
        zahlenformate={
            "anzahl": "#,##0",
            "mittelwert_sekunden": "0.##",
            "median_sekunden": "0.##",
        },
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Potenzielle Wartezeit · Gl. 3.15")
    zeile = _paare(
        ws,
        zeile,
        (
            (
                "Hinweis",
                "Positive Zeitdifferenzen werden als potenzielle Wartezeiten ausgewiesen. "
                "Sie stellen ohne fachliche Bestätigung keine Warteschlange dar.",
            ),
        ),
    )
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("Übergang", "uebergang"),
            ("n", "anzahl"),
            ("Mittelwert (s)", "mittelwert_sekunden"),
            ("Median (s)", "median_sekunden"),
            ("Status", "status"),
        ),
        [
            dict(wert)
            for wert in _liste(zeitwahl.get("potenzielle_wartezeiten_tabelle"))
            if isinstance(wert, Mapping)
        ],
        autofilter=True,
        zahlenformate={
            "anzahl": "#,##0",
            "mittelwert_sekunden": "0.##",
            "median_sekunden": "0.##",
        },
    )
    vereinfachte = [
        dict(wert)
        for wert in _liste(zeitwahl.get("vereinfachte_zeitspannen_tabelle"))
        if isinstance(wert, Mapping)
    ]
    if vereinfachte:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Vereinfachte Zeitspannen")
        zeile = _paare(
            ws,
            zeile,
            (
                (
                    "Einordnung",
                    "Bestätigte Start-zu-Start-Zeitspannen enthalten mögliche "
                    "Bearbeitungs-, Transport- und Warteanteile gemeinsam.",
                ),
            ),
        )
        _tabelle(
            ws,
            zeile,
            (
                ("Übergang", "uebergang"),
                ("n", "anzahl"),
                ("Mittelwert (s)", "mittelwert_sekunden"),
                ("Median (s)", "median_sekunden"),
            ),
            vereinfachte,
            autofilter=True,
            zahlenformate={
                "anzahl": "#,##0",
                "mittelwert_sekunden": "0.##",
                "median_sekunden": "0.##",
            },
        )


def _analyse(ws: Worksheet, report: Mapping[str, Any]) -> None:
    prozess = _mapping(report.get("prozessdarstellung"))
    ausgaben = _mapping(report.get("ausgaben"))
    ressourcen = _mapping(report.get("ressourcen"))
    warteschlangen = _mapping(report.get("warteschlangen"))
    zeile = _sheet_start(
        ws, SHEET_NAMES[7], "Process Mining, Conformance, Zeit- und Ressourcenbefunde"
    )
    zeile = _abschnitt(ws, zeile, "Process Discovery und Conformance")
    zeile = _paare(
        ws,
        zeile,
        (("Notation", prozess.get("notation_anzeige")),),
    )
    conformance = _mapping(ausgaben.get("conformance_checking"))
    conformance_ergebnis = _mapping(conformance.get("ergebnis"))
    if conformance_ergebnis:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Token-Based Replay · Gleichung 3.14")
        zeile = _paare(
            ws,
            zeile,
            (
                (
                    "Fitness",
                    conformance_ergebnis.get(
                        "fitness_anzeige", conformance_ergebnis.get("fitness")
                    ),
                ),
                ("Produzierte Tokens pT", conformance_ergebnis.get("produzierte_tokens")),
                ("Konsumierte Tokens cT", conformance_ergebnis.get("konsumierte_tokens")),
                ("Fehlende Tokens mT", conformance_ergebnis.get("fehlende_tokens")),
                ("Verbleibende Tokens rT", conformance_ergebnis.get("verbleibende_tokens")),
                ("Ausgewertete Fälle", conformance_ergebnis.get("ausgewertete_faelle")),
                ("Konforme Fälle", conformance_ergebnis.get("konforme_faelle")),
                ("Abweichende Fälle", conformance_ergebnis.get("abweichende_faelle")),
            ),
        )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Kennzahlen- und Ressourcenbefunde")
    kpi_zeilen = [
        dict(wert)
        for wert in (
            _liste(ausgaben.get("kpi_ergebnisse"))
            + _liste(warteschlangen.get("wartezeit_kpis"))
            + _liste(ressourcen.get("ressourcenbezogene_kpis"))
        )
        if isinstance(wert, Mapping)
    ]
    for kpi in kpi_zeilen:
        if not isinstance(kpi.get("ergebnis"), (int, float)):
            kpi["ergebnis"] = kpi.get("ergebnis_anzeige")
    zeile = _tabelle(
        ws,
        zeile,
        (
            ("KPI", "bezeichnung"),
            ("Ergebnis", "ergebnis"),
            ("Formel", "formel"),
            ("Bezugsmenge", "bezugsmenge"),
            ("Status", "status_anzeige"),
        ),
        kpi_zeilen,
        autofilter=True,
        zahlenformate={"ergebnis": "0.##"},
    )
    performance = _mapping(ausgaben.get("performance_und_engpassanalyse"))
    dt_db = _mapping(performance.get("dt_db_ergebnis"))
    performance_zeilen = []
    for art, schluessel in (
        ("dT · Fertigstellungsabweichung", "dt_statistik"),
        ("dB · Bearbeitungszeitabweichung", "db_statistik"),
    ):
        statistik = _mapping(dt_db.get(schluessel))
        if statistik:
            performance_zeilen.append({"art": art, **statistik})
    if performance_zeilen:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Soll-/Ist-Abweichungen")
        zeile = _tabelle(
            ws,
            zeile,
            (
                ("Analyse", "art"),
                ("n", "anzahl"),
                ("Mittelwert (s)", "mittelwert_sekunden"),
                ("Median (s)", "median_sekunden"),
                ("Verspätet", "verspaetet"),
                ("Planmäßig", "planmaessig"),
                ("Vorzeitig", "vorzeitig"),
                ("Länger", "laenger_als_geplant"),
                ("Gleich", "gleich_geplant"),
                ("Kürzer", "kuerzer_als_geplant"),
            ),
            performance_zeilen,
            zahlenformate={"mittelwert_sekunden": "0.##", "median_sekunden": "0.##"},
        )
    busy = _mapping(performance.get("busy_ratio_ergebnis"))
    busy_zeilen = [
        dict(wert)
        for wert in _liste(busy.get("ressourcenstatistiken"))
        if isinstance(wert, Mapping)
    ]
    if busy_zeilen:
        zeile += 1
        zeile = _abschnitt(ws, zeile, "Ergänzende Performance · Busy Ratio")
        zeile = _tabelle(
            ws,
            zeile,
            (
                ("Ressource", "ressource"),
                ("n", "anzahl_gueltige_busy_ratios"),
                ("Mittelwert", "mittelwert_busy_ratio"),
                ("Median", "median_busy_ratio"),
                ("Minimum", "minimum_busy_ratio"),
                ("Maximum", "maximum_busy_ratio"),
            ),
            busy_zeilen,
        )
        zeile = _paare(
            ws,
            zeile,
            (
                ("Potenzieller Engpass", busy.get("potenzieller_engpass")),
                (
                    "Einordnung",
                    "Aus der Busy Ratio wird keine Ursache und keine reale "
                    "Warteschlange abgeleitet.",
                ),
            ),
        )
        busy_einzelwerte = [
            dict(wert) for wert in _liste(busy.get("einzelwerte")) if isinstance(wert, Mapping)
        ]
        if busy_einzelwerte:
            zeile += 1
            zeile = _abschnitt(
                ws,
                zeile,
                "Busy-Ratio-Details · Zwischenankunftszeit Ressource · Gl. 3.4",
            )
            zeile = _tabelle(
                ws,
                zeile,
                (
                    ("Ressource", "ressource"),
                    ("Aktivität", "aktivitaet"),
                    ("Bearbeitungszeit (s)", "bearbeitungszeit_sekunden"),
                    (
                        "Zwischenankunftszeit Ressource (s)",
                        "ressourcenbezogene_zwischenankunftszeit_sekunden",
                    ),
                    ("Busy Ratio", "busy_ratio"),
                ),
                busy_einzelwerte,
                autofilter=True,
                zahlenformate={
                    "bearbeitungszeit_sekunden": "0.##",
                    "ressourcenbezogene_zwischenankunftszeit_sekunden": "0.##",
                    "busy_ratio": "0.####",
                },
            )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Aktivitäts-/Ressourcenzuordnung")
    _tabelle(
        ws,
        zeile,
        (
            ("Aktivität", "aktivitaet"),
            ("Ressourcen", "ressourcen"),
            ("Modus", "modus"),
            ("Herkunft", "herkunft"),
        ),
        [
            dict(wert)
            for wert in _liste(ressourcen.get("aktivitaet_ressourcen"))
            if isinstance(wert, Mapping)
        ],
    )


def _validierung(ws: Worksheet, report: Mapping[str, Any]) -> None:
    validierung = _mapping(report.get("validierung"))
    zeile = _sheet_start(
        ws, SHEET_NAMES[8], "Vorschlag, finale Fassung und Validierungsentscheidung"
    )
    zeile = _abschnitt(ws, zeile, "Gesamtvalidierung")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Status", validierung.get("status_anzeige")),
            ("Validierungsvermerk", validierung.get("validierungsvermerk")),
            ("Menschlich bestätigt", validierung.get("menschlich_bestaetigt")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Validierte Modellbestandteile")
    vergleich = []
    for bestandteil in _liste(report.get("modellbestandteile")):
        if not isinstance(bestandteil, Mapping):
            continue
        anpassungen = _liste(bestandteil.get("fachliche_anpassungen"))
        entscheidungen = _liste(bestandteil.get("fachliche_entscheidungen"))
        vergleich.append(
            {
                "bestandteil": bestandteil.get("bezeichnung"),
                "entscheidungen": len(entscheidungen),
                "anpassungen": len(anpassungen),
                "geaendert": bool(anpassungen),
                "status": bestandteil.get("validierungsstatus_anzeige"),
            }
        )
    _tabelle(
        ws,
        zeile,
        (
            ("Bestandteil", "bestandteil"),
            ("Fachliche Entscheidungen", "entscheidungen"),
            ("Ergänzungen/Anpassungen", "anpassungen"),
            ("Geändert", "geaendert"),
            ("Status", "status"),
        ),
        vergleich,
        autofilter=True,
        zahlenformate={"entscheidungen": "#,##0", "anpassungen": "#,##0"},
    )


def _lineage(ws: Worksheet, report: Mapping[str, Any]) -> None:
    projekt = _mapping(report.get("projekt"))
    modell = _mapping(report.get("modell"))
    prozess = _mapping(report.get("prozessdarstellung"))
    lineage = _mapping(report.get("lineage"))
    dokument = _mapping(report.get("dokument"))
    zeile = _sheet_start(ws, SHEET_NAMES[9], "IDs, Referenzen, Prüfsummen und Artefaktherkunft")
    zeile = _abschnitt(ws, zeile, "Modellreferenzen")
    zeile = _paare(
        ws,
        zeile,
        (
            ("Projekt-ID", projekt.get("projekt_id")),
            ("K*-ID", modell.get("k_stern_id")),
            ("Validierungslauf-ID", modell.get("validierungslauf_id")),
            ("Prozessmodell-ID", prozess.get("prozessmodell_id")),
            ("Analyse-ID", prozess.get("process_mining_analyse_id")),
            ("K-Referenz", lineage.get("k_referenz")),
            ("O-Referenz", lineage.get("o_referenz")),
            ("Eingabefingerabdruck", lineage.get("eingabefingerabdruck")),
            ("Entscheidungsfingerabdruck", lineage.get("entscheidungsfingerabdruck")),
            ("Gesamtprüfsumme", lineage.get("gesamtpruefsumme")),
            ("Framework-Version", dokument.get("softwareversion")),
        ),
    )
    zeile += 1
    zeile = _abschnitt(ws, zeile, "Informationsherkunft")
    _tabelle(
        ws,
        zeile,
        (
            ("Bestandteil", "bestandteil_id"),
            ("Informations-ID", "informations_id"),
            ("Strukturreferenz", "strukturreferenz"),
            ("Herkunftsartefakt", "herkunftsartefakt"),
            ("Artefakt-ID", "herkunftsartefakt_id"),
            ("Prüfsumme", "herkunftsartefakt_sha256"),
            ("Übernahmeart", "uebernahmeart_anzeige"),
        ),
        [dict(wert) for wert in _liste(lineage.get("informationen")) if isinstance(wert, Mapping)],
        autofilter=True,
    )


def _arbeitsmappe_finalisieren(arbeitsmappe: Workbook) -> None:
    breite_grenzen: dict[str, dict[str, float]] = {
        SHEET_NAMES[0]: {"B": 72},
        SHEET_NAMES[1]: {"B": 82, "D": 48},
        SHEET_NAMES[2]: {"B": 48, "E": 34, "F": 62},
        SHEET_NAMES[4]: {"B": 42, "D": 62, "E": 42, "F": 54},
        SHEET_NAMES[5]: {"C": 62, "D": 62},
        SHEET_NAMES[6]: {"A": 32, "B": 48, "G": 45},
        SHEET_NAMES[7]: {"A": 38, "C": 55, "D": 42},
        SHEET_NAMES[8]: {"B": 62, "C": 62, "D": 54},
        SHEET_NAMES[9]: {"C": 54, "D": 38, "E": 42, "F": 45},
    }
    for ws in arbeitsmappe.worksheets:
        if ws.sheet_properties.pageSetUpPr is not None:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = False
        if ws.max_column > 5:
            ws.page_setup.orientation = "landscape"
        for spalte in range(1, max(8, ws.max_column) + 1):
            buchstabe = get_column_letter(spalte)
            if ws.title == SHEET_NAMES[3]:
                continue
            inhaltslaengen = []
            for zelle in (ws.cell(zeile, spalte) for zeile in range(4, ws.max_row + 1)):
                if zelle.value is None or zelle.coordinate in ws.merged_cells:
                    continue
                inhaltslaengen.extend(len(teil) for teil in str(zelle.value).splitlines())
            minimum = 14.0 if spalte > 1 else 20.0
            maximum = breite_grenzen.get(ws.title, {}).get(buchstabe, 36.0)
            passend = max(inhaltslaengen, default=int(minimum)) + 2
            ws.column_dimensions[buchstabe].width = min(maximum, max(minimum, passend))
        for zeile in ws.iter_rows():
            for zelle in zeile:
                if zelle.value is not None:
                    bestehend = zelle.font
                    zelle.font = Font(
                        name=FONT_FAMILY,
                        size=bestehend.sz or 12,
                        bold=bestehend.bold,
                        italic=bestehend.italic,
                        color=bestehend.color,
                        underline=bestehend.underline,
                    )
                if zelle.value is not None and zelle.alignment == Alignment():
                    zelle.alignment = Alignment(wrap_text=True, vertical="top")


def render_report_xlsx(report_data: Mapping[str, Any]) -> bytes:
    """Rendert die gemeinsamen, bereits aufgelösten Reportdaten als XLSX-Bytes."""
    if report_data.get("report_data_version") != REPORT_DATA_VERSION:
        raise XlsxRenderingFehler(
            f"Nicht unterstützte Report-Datenversion: {report_data.get('report_data_version')!r}."
        )
    try:
        arbeitsmappe = Workbook()
        standardblatt = arbeitsmappe.active
        if standardblatt is not None:
            arbeitsmappe.remove(standardblatt)
        funktionen = (
            _uebersicht,
            _problem_und_grenze,
            _ziele,
            _prozess,
            _systemelemente,
            _annahmen,
            _daten,
            _analyse,
            _validierung,
            _lineage,
        )
        for name, funktion in zip(SHEET_NAMES, funktionen, strict=True):
            funktion(arbeitsmappe.create_sheet(name), report_data)
        _arbeitsmappe_finalisieren(arbeitsmappe)
        ziel = BytesIO()
        arbeitsmappe.save(ziel)
        return ziel.getvalue()
    except XlsxRenderingFehler:
        raise
    except (OSError, TypeError, ValueError, KeyError) as fehler:
        raise XlsxRenderingFehler(
            f"Die XLSX-Arbeitsmappe konnte nicht erzeugt werden: {fehler}"
        ) from fehler
